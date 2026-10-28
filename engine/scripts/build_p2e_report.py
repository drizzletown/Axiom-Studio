from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
from datetime import date

ROOT=Path(__file__).resolve().parents[1]
out=ROOT/'docs'/'Axiom_P2_第五批实施与验收报告_V0.5_20261025.docx'
doc=Document()
sec=doc.sections[0]
sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2.0)
sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.3)
sec.header_distance=Cm(0.9);sec.footer_distance=Cm(1)
styles=doc.styles
normal=styles['Normal'];normal.font.name='Noto Sans CJK SC';normal.font.size=Pt(10)
normal.font.color.rgb=RGBColor(35,44,55)
normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.2
for name,sz in [('Title',18),('Heading 1',13),('Heading 2',11)]:
    s=styles[name];s.font.name='Noto Sans CJK SC';s.font.size=Pt(sz);s.font.bold=True;s.font.color.rgb=RGBColor(30,57,87)
    s.paragraph_format.space_before=Pt(9);s.paragraph_format.space_after=Pt(4)
header=sec.header.paragraphs[0];header.text='AXIOM  |  ENGINEERING VERIFICATION SERIES';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor(100,112,125)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
foot.add_run('Axiom 0.2 冻结基线  ·  工程原型 V0.7.0  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)

def p(text='',style=None):return doc.add_paragraph(text,style)
def heading(t,l=1):doc.add_heading(t,l)
def table(cols,rows):
    t=doc.add_table(rows=1, cols=len(cols));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,c in enumerate(cols):t.rows[0].cells[i].text=c
    for item in rows:
        cells=t.add_row().cells
        for i,value in enumerate(item):cells[i].text=str(value)
    for c in t.rows[0].cells:
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcPr=c._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EAF1F8');tcPr.append(shd)
        for run in c.paragraphs[0].runs:run.bold=True
    for row in t.rows:
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0)
                for run in para.runs:run.font.size=Pt(8.5)
    doc.add_paragraph('')

s=doc.add_paragraph(style='Title');s.alignment=WD_ALIGN_PARAGRAPH.CENTER;s.add_run('Axiom P2 第五批实施与验收报告')
s=doc.add_paragraph();s.alignment=WD_ALIGN_PARAGRAPH.CENTER;s.add_run('版本 V0.5  |  工程版本 V0.7.0  |  2026-10-25')
p('文档属性：阶段验收记录；对应《Axiom 概要设计说明书 V1.0》与 Axiom 0.2 冻结基线。')
heading('一、阶段结论')
p('本轮实现 Linux 用户／挂载／网络／PID 命名空间中的受限 C++ 编译与执行，并将真实 C++ 双组件候选与 Axiom 有限交错行为模型连接，执行多步观测轨迹差分检查。P0—P2 的历史功能继续保留。')
p('验收结果：134 项自动化测试通过；深度 4 的行为轨迹测试对正确 C++ 候选给出 Tested，针对单侧状态更新及未复位隐藏状态的缺陷候选给出 Violated；本地模型＋C++ 测试级 Gate 按预期阻断或批准。')
p('重要边界：本轮“Approved”仅指 cpp-deterministic-schedule-tested-prototype；release_authorized 始终为 false。既不表示实际多线程 C++ 正确，也不能作为生产运行不可信代码的安全认证。')
heading('二、本轮实施范围')
table(['模块','交付内容','状态'],[
    ('Linux Namespace Runner','user/mount/network/pid 隔离、只读 /usr、chroot 工作区、进程超时及输出限制','完成受限原型'),
    ('Cpp Trace Adapter','Source/Mirror Bool ABI；set 同步、peek 独立操作；每个场景执行 reset','完成受限原型'),
    ('差分与反例','生成所有长度不超过 D 的 API 原子调用序列；逐步检查观察值和后继状态','完成'),
    ('候选质量检查','正常与 UBSan 双编译变体；故障、失配、超时不放行','完成'),
    ('证据与 Gate','源码／模型／Manifest／工具摘要绑定、测试证据本地审计与过期阻断','完成'),
])
heading('三、语义与检测范围')
p('本轮测试专门针对两个布尔状态组件，采用固定 ABI：axiom_reset(left,right) 与 axiom_step(opcode,desired,&left,&right)。验证器在每个场景中依次执行 C++ 调用，并与 Axiom 的抽象后继集合比较；随机调度、指令级竞态、内存模型、故障恢复及公平性均未纳入验证。')
p('当 D=4 时，存在 4 种原子调用选择，每种序列双向重放。共执行 680 条调度序列（正常与 UBSan 各执行一次），每种变体包含 2504 个操作观察。固定调度穷举不能外推为所有 C++ 并发执行正确。')
doc.add_page_break()
heading('四、验收记录')
table(['检查对象','预期','实测结果'],[
    ('正确候选 correct.cpp','Tested；满足测试级 Gate','Tested / Approved / release=false'),
    ('broken_one_side.cpp','发现跨组件状态偏差','Violated；输出 [1,0,1]'),
    ('broken_hidden_state.cpp','发现跨场景隐藏状态污染','Violated；生成反例'),
    ('旧模型证据缺失','阻断 C++ 测试级 Gate','Blocked'),
    ('源码修改或深度改变','不复用旧候选证据','Blocked'),
    ('证据对象被篡改','审计发现问题并阻断','Findings / Blocked'),
    ('Linux 隔离探测','隔离可用才执行；不可用阻断','当前测试环境可用'),
    ('自动化测试','P0—P2e 全部回归','134/134 通过'),
])
heading('五、安全边界与剩余风险')
p('现有命名空间构成一定程度的进程、网络和文件系统隔离，但并未实施完备的 seccomp syscall 白名单、独立宿主机／虚拟机隔离、强制 cgroup 内存上限、内核安全审计与恶意代码对抗测试；不能声称可安全运行任意敌意 C++。传统 P1 cpp-verify 仍然要求 --trusted-local。')
p('代码实现仍可能存在未枚举到的调度行为、编译器／标准库特有行为、实际线程层面的数据竞争，以及规格本身的缺陷。离线签名只认证指定验证记录，不自动提升真实 C++ 语义保证等级。')
heading('六、下一阶段工作')
p('下一批优先将隔离执行路径进一步硬化并纳入可信构建环境，开展真实线程调度与线性化历史采集；随后在更复杂的并发账本／库存项目上进行连续版本的实现级增量验证实验，衡量缺陷检出率、漏检风险与重复验证成本。所有新结果继续保持证明、有限模型、有界测试与运行观测的严格分类。')
heading('附录：复现实验')
p('python -m pip install -e ".[test]"')
p('python -m pytest -q')
p('python scripts/run_p2e_demo.py')
p('ax --workspace .ax-p2e cpp-trace-verify examples/interleaving/product_safe.json --candidate cpp/trace/correct.cpp --depth 4')
p('完整工程交付包内含 README、cpp/trace 示例、tests/test_cpp_trace_p2e.py、reports/p2e_demo.json 和 pytest 测试记录。')
doc.save(out)
print(out)
