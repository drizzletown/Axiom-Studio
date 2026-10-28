"""Create the P2d acceptance report from validated run results."""
from __future__ import annotations
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs'/'Axiom_P2_第四批实施与验收报告_V0.4_20261025.docx'
actual=json.loads((ROOT/'reports/p2d_demo.json').read_text(encoding='utf8'))
assert actual['safe']['outcome']=='Bounded' and actual['broken']['outcome']=='Violated'

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);tcPr.append(shd)

def set_cell(cell,value,bold=False):
    cell.text=str(value);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_after=Pt(0)
        for r in p.runs:r.bold=bold;r.font.size=Pt(8.45)

def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1, cols=len(headers));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):set_cell(t.rows[0].cells[i],h,True);shade(t.rows[0].cells[i],'E9EFF5')
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):set_cell(cells[i],v)
    if widths:
        for row in t.rows:
            for c,w in zip(row.cells,widths):c.width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def para(doc,txt,style=None):
    return doc.add_paragraph(txt,style=style)

def heading(doc,txt,level=1):return doc.add_heading(txt,level)

def add_code(doc,txt):
    for line in txt.strip().splitlines():
        p=doc.add_paragraph(style='Normal');p.paragraph_format.left_indent=Inches(.23)
        p.paragraph_format.space_after=Pt(1)
        run=p.add_run(line);run.font.name='Consolas';run.font.size=Pt(8);run.font.color.rgb=RGBColor(55,63,72)

D=Document();sect=D.sections[0]
sect.top_margin=Inches(.65);sect.bottom_margin=Inches(.62)
sect.left_margin=Inches(.83);sect.right_margin=Inches(.77)
sect.header_distance=Inches(.35);sect.footer_distance=Inches(.37)
styles=D.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    s=styles[name]
    s.font.name='Noto Sans CJK SC'
    s._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
styles['Normal'].font.size=Pt(9.2);styles['Normal'].paragraph_format.space_after=Pt(3)
styles['Normal'].paragraph_format.line_spacing=1.1
for name,size in [('Title',18),('Heading 1',12),('Heading 2',10)]:
    styles[name].font.size=Pt(size)
    styles[name].font.color.rgb=RGBColor(27,51,75)
    styles[name].font.bold=True
    styles[name].paragraph_format.space_before=Pt(9)
    styles[name].paragraph_format.space_after=Pt(4)

header=sect.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('AXIOM  |  P2  IMPLEMENTATION REVIEW');r.font.size=Pt(8);r.font.color.rgb=RGBColor(100,111,121)
footer=sect.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Axiom 0.2 基线不变  ·  概要设计 V1.0  |  第 ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
footer.add_run(' 页').font.size=Pt(8)

p=D.add_paragraph(style='Title');p.add_run('Axiom P2 第四批实施与验收报告')
p=D.add_paragraph();p.style='Subtitle';p.add_run('V0.4  ·  工程版本 0.6.0  ·  2026 年 10 月 25 日')
para(D,'基于 Axiom 0.2 冻结理论基线与《Axiom 概要设计说明书 V1.0》。本版继续原有 P0—P2c 工程，新增混合同步/交错行为模型检查，以及独立签名的有限抽象验证证据机制。')
heading(D,'1 交付结论')
table(D,['验收项目','实测结果','结论'],[
 ('原有 P0—P2c 自动化回归','96 项原有测试通过','通过'),
 ('本轮新增自动化测试','25 项新增测试通过','通过'),
 ('合并自动化测试','121 / 121','通过'),
 ('混合同步/交错组合验证','正确模型 Bounded；错误模型 Violated','通过'),
 ('签名报告验真与失效阻断','Ed25519 验签、重放、数据绑定测试通过','通过'),
 ('不可信 C++ 安全隔离','尚无经审计的 OS 执行沙箱','未完成'),
 ('真实 C++ 并发精化验证','尚未建立受信任的实现级行为绑定','未完成')],widths=[2.05,3.6,1.1])
para(D,'验收判定：P2 第四批受限原型通过；完整 P2 与生产发布准入尚未完成。任何 Bounded / Approved 均限定于明确的有限抽象模型与指定证据条件。')
heading(D,'2 设计内容与体系边界')
heading(D,'2.1 混合同步/交错验证',2)
para(D,'新增 ax-interleaving-composition/0.1。两组件的 Action 必须且只能归入两类：原子同步对（sync），或独立原子操作（independent）。检查器采用广度优先穷举，覆盖所有可达状态并在每个全局状态验证关系不变量。对于同步对，两侧共享同一输入并要求输出相等。')
para(D,'该机制能发现同步产品无法表现的单边更新错误。但不支持线程内部指令级交错、非原子调用、网络消息、故障恢复、公平性或活性结论。')
heading(D,'2.2 证据准入与签名',2)
para(D,'新增 interleave-verify / interleave-gate 以及独立的 Ed25519 签发与验签接口。Gate 要求组件级模型证据与当前绑定版本的组合证据同时存在。签名者在独立重放本地检查器并通过工件审计后，对证据及绑定摘要进行签名；验收者使用事先信任的公钥和 signer ID 复验。')
para(D,'签名提高了对工作区内容被篡改的识别能力，但其保证取决于独立受保护的签名者、公钥分发和可信验签程序。该机制不能证明 C++ 代码正确，也不能将本地运行环境变为安全沙箱。')

heading(D,'3 关键实验及反例')
table(D,['场景','观察结果'],[
 ('受控同步修改 + 独立读取','Bounded；可达全局状态 2 个；候选检查 48 次'),
 ('允许两端分别独立修改','Violated；违反 MirroredState'),
 ('缺少组件独立证据','Gate = Blocked'),
 ('补齐组件证据及组合证据','有限抽象模型 Gate = Approved'),
 ('证据/签名/模型被修改','验签失败或准入阻断'),
 ('资源预算不足','Unknown，不得准入')],widths=[2.5,4.25])
D.add_page_break()
para(D,'错误组合反例（最短全局动作轨迹）：')
trace=actual['broken']['counterexample']
add_code(D,f'''S0: Source.on = {str(trace[0]['state']['Source.on']).lower()}, Mirror.on = {str(trace[0]['state']['Mirror.on']).lower()}
S1: {trace[1]['action']}(desired = true)
S1: Source.on = false, Mirror.on = true
Violation: Source.on == Mirror.on is false''')
para(D,'两个组件各自满足本地 WellTyped 性质，但在允许独立执行 set 的全局调度下违反跨组件不变量，表明局部模型证据不可替代组合行为检查。')
heading(D,'4 代码模块及接口')
table(D,['模块/命令','职责'],[
 ('axiom.interleaving','schema 与绑定校验；有限同步/独立交错状态探索；反例生成'),
 ('axiom.interleaving_store','当前证据索引、独立重放、模型级 Gate 和失效控制'),
 ('axiom.attestation','Ed25519 离线签名、绑定证明、使用固定公钥验签'),
 ('axiom.audit','新增 P2d 工件一致性审计'),
 ('ax interleave-verify / interleave-gate','有限模型验证与准入控制'),
 ('ax interleave-attest / interleave-verify-attestation','受信任签名与独立验签')],widths=[3.05,3.7])
heading(D,'5 安全审查与剩余风险')
para(D,'（1）不可信执行：当前机器未完成经审计的 OS 沙箱，cpp-verify 继续要求 --trusted-local。进程资源限额不等于沙箱。')
para(D,'（2）证据可信性：独立签名须在与候选生成隔离的权限域执行，私钥不得保存在工作区。验签方必须固定公钥，避免被替换。')
para(D,'（3）模型充分性：当前静态有限模型仅能验证已表达的安全性质，不足以推断业务需求完整、环境正确或真实系统满足规格。')
para(D,'（4）实现精化：P2d 并未证明实际 C++ 对抽象状态机的实现精化。C++ 验证仍限先前 P1 的测试级机制。')
heading(D,'6 后续阶段准入条件')
para(D,'优先事项 A：采用受控签名/验签运行环境，实现由独立服务承担的质量准入，不与 AI 候选生成进程共享签名密钥。')
para(D,'优先事项 B：实现可信 OS 沙箱，支持受控编译和执行不可信 C++；在沙箱完成安全审计前不得扩大现有执行权限。')
para(D,'优先事项 C：建立真实 C++ 并发轨迹与抽象 Action/Observation 的映射，并与模型进行有范围的差分验证；不得将测试证据升级为通用形式化证明。')
para(D,'优先事项 D：以真实多轮代码修改比较增量与全量验证的成本、准确性及证据复用风险，为 P2 验收及 P3 准入提供数据。')
heading(D,'7 版本一致性')
para(D,'Axiom 0.2 理论基线 AX-B01—AX-B09 未变；概要设计 V1.0 未变；P0—P2c 已有接口保留。新功能全部使用明确版本号与独立工件，不将历史测试成果伪称为新的实现级形式化保证。')
D.save(OUT)
print(OUT)
