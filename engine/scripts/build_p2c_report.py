"""Produce the P2c phase report from reproduced program results."""
from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
import re

ROOT=Path(__file__).resolve().parents[1]
report=json.loads((ROOT/'reports/p2c_demo.json').read_text())
test=(ROOT/'reports/pytest_p2c.txt').read_text()
filename=ROOT/'docs/Axiom_P2_第三批实施与验收报告_V0.3_20261025.docx'
doc=Document();sec=doc.sections[0]
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.9)
sec.left_margin=Cm(2.55);sec.right_margin=Cm(2.45)

styles=doc.styles
for n in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
    s=styles[n]
    s.font.name='Noto Sans CJK SC'
    s._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
styles['Normal'].font.size=Pt(10)
styles['Normal'].paragraph_format.space_after=Pt(5)
styles['Normal'].paragraph_format.line_spacing=1.14
styles['Title'].font.size=Pt(21)
styles['Title'].font.color.rgb=RGBColor(0,0,0)
styles['Subtitle'].font.size=Pt(10.5)
styles['Subtitle'].font.color.rgb=RGBColor(0,0,0)
styles['Subtitle'].font.italic=False
for n,sz in (('Heading 1',15),('Heading 2',12),('Heading 3',11)):
    styles[n].font.size=Pt(sz)
    styles[n].font.color.rgb=RGBColor(0,0,0)
    styles[n].paragraph_format.space_before=Pt(8)
    styles[n].paragraph_format.keep_with_next=True
    styles[n].paragraph_format.space_after=Pt(6)

header=sec.header.paragraphs[0];header.text='AXIOM · 工程实施文档                                    P2c / 0.5.0'
header.style=styles['Normal'];header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor(95,95,95)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Axiom 0.2 冻结理论基线   |   HLD 1.0   |   ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def p(txt,style=None): return doc.add_paragraph(txt,style=style)
def h(txt,lev=1):doc.add_heading(txt,level=lev)
def t(headers,rows):
    table=doc.add_table(rows=1,cols=len(headers));table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(headers):table.rows[0].cells[i].text=str(x)
    for row in rows:
        cells=table.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    for i,row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(1)
                for r in par.runs:r.font.size=Pt(9)
            if i==0:
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EAEAEA');cell._tc.get_or_add_tcPr().append(shd)
    doc.add_paragraph('').paragraph_format.space_after=Pt(0)

p('Axiom P2 第三批实施与验收报告', 'Title')
p('工程版本 0.5.0  ·  报告 V0.3  ·  2026 年 10 月 25 日','Subtitle')
p('适用范围：双组件有限抽象模型的同步行为组合验证；继承 P0、P1、P2a、P2b 工程成果。')

h('1  项目基线与阶段结论')
p('依据《Axiom 概要设计说明书 V1.0》与 AX-B01—AX-B09 冻结条款推进。本阶段解决 P2b 仅能检查组件间显式证据依赖、无法检查组合后可观察行为的问题。')
p('阶段结论：P2c 的受限同步抽象乘积检查与证据准入闭环已经实现并通过回归测试；完整 P2 尚未验收。模型结论不能代替真实 C++ 程序的实现级精化证明。')
t(['评价项','已实现的结果'],[
    ('抽象组合','两个有限模型，全 Action 显式同步配对、相同输入域、输出相等'),
    ('检查对象','乘积可达状态、跨组件关系不变量、同步输出条件'),
    ('报告语义','Bounded / Violated / Unknown，失败时包含同步步骤反例'),
    ('质量准入','两个局部模型证据 + 同版本组合证据缺一不可'),
    ('工件完整性','状态模型、同步协议、证据报告和工具代码摘要绑定；本地审计'),
    ('不覆盖的内容','任意交错、多组件网络执行、C++ 组合精化、不可信代码沙箱')])

h('2  本轮开发成果')
t(['模块/工件','新增职责'],[
    ('src/axiom/product.py','规范化 manifest、动作配对验证、类型化关系性质、同步乘积 BFS 穷举及反例生成'),
    ('src/axiom/product_store.py','Bounded 报告重放、内容寻址证据、同步组合 Gate'),
    ('src/axiom/audit.py','扩充产品证据索引、校验工具摘要及内容寻址对象一致性'),
    ('src/axiom/cli.py','product-verify、product-gate 命令'),
    ('examples/product/','正确镜像、状态缺陷、输出缺陷样例及 manifest'),
    ('tests/test_product_p2c.py','属性违反、输出冲突、资源上限、版本失效、伪造报告、路径攻击等回归测试')])

h('3  关键验证结果')
p('完整自动化回归：'+re.search(r'\d+ passed in [\d.]+s',test).group(0)+'。原有 P1 的真实 C++ 演示仍正常执行，故障候选中的错误映射、边界错误、隐藏状态和整数溢出均按预期被检出。')
t(['验证对象','结果','可支持的结论'],[
    ('Source 局部模型','Bounded','局部状态性质在其有限抽象域内得到穷举检查'),
    ('Mirror 局部模型','Bounded','局部状态性质在其有限抽象域内得到穷举检查'),
    ('正确同步组合','Bounded；可达乘积状态 2','关系不变量与输出要求在给定同步组合模型内满足'),
    ('错误 Mirror 的局部模型','Bounded','单独看两个模块无法暴露关系错误'),
    ('错误同步组合','Violated；E_PRODUCT_PROPERTY','得到一次同步动作后的组合反例'),
    ('证据篡改或规格变动','Blocked','过期或损坏证据不会被组合 Gate 直接使用'),
    ('本地完整性审计','Consistent；2 条组合运行记录','目前保存的本地工件与索引一致，不构成恶意攻击安全保证')])

h('4  关键反例与工程意义')
p('初始 Source.on=false、Mirror.on=false；共同执行 set(desired=false)，两个 Action 的输出均为 true。正确组件的后继状态保持 false，错误 Mirror 的后继状态变为 true，导致 Source.on == Mirror.on 被违反。')
p('这验证了 P2c 相比 P2b 的增量能力：两个独立模型全部通过局部检查，仍然可能存在跨组件行为不一致；同步乘积关系检查可在明确的语义范围内识别这种问题。')

h('5  可信边界与风险')
p('本版只检查双模型原子同步步骤。未配对 Action 一律拒绝；任意交错、异步消息、失败重试与环境扰动未建模。有限穷举的 Bounded 是模型级结论。')
p('证据由本地内容寻址文件与 SQLite 保存，能够发现一般内容篡改及过期绑定；具有相同工作区写权限的攻击者仍可能伪造全套一致内容，故尚不具备独立信任根。')
p('真实 C++ 执行仍需 --trusted-local，未提供经过审计的隔离沙箱。当前 Gate 的 Approved 仅表示 finite-abstract-product-integration-check，绝不代表生产发布许可。')

h('6  下一阶段顺序')
p('（1）优先完成不可信代码安全执行隔离、独立证据签名与可信验证器部署；（2）扩展受控交错与三组件以上组合语义；（3）明确 C++ 接口/状态抽象映射及实现级证据；（4）对连续软件演化开展完整验证与增量验证的成本及漏检风险对照实验。')
p('Axiom 0.2 理论基线继续冻结，后续调整仅限概要设计允许的工程实现层；任何理论内核变更须单独提出。')

filename.parent.mkdir(parents=True,exist_ok=True)
doc.save(filename)
print(filename)
